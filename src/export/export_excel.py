from __future__ import annotations
from typing import List, Dict, Optional
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference

from simulation.particle_motion import Particle
from simulation.rtd import (
    residence_time_summary,
    collect_residence_times,
    compute_exit_count_histogram,
    compute_cumulative_exit_counts_excel,
    mean_residence_time_by_group,
    count_active_particles_by_reactor,
)


# Construction des DataFrames (une fonction par feuille)
def _particles_to_dataframe(particles: List[Particle]) -> pd.DataFrame:
    rows = []
    for i, p in enumerate(particles):
        rows.append({
            "id": i + 1,
            "densite_kg_m3": p.particle_type.particle_density,
            "rayon_mm": p.particle_type.particle_size * 1000.0,
            "reacteur_final": p.current_reactor,
            "temps_entree_min": p.entry_time_s / 60.0 if p.entry_time_s is not None else None,
            "temps_sortie_min": p.exit_time_s / 60.0 if p.exit_time_s is not None else None,
            "temps_residence_min": p.residence_time_min,
            "est_sortie": p.residence_time_min is not None,
        })
    return pd.DataFrame(rows)


def _summary_to_dataframe(particles: List[Particle]) -> pd.DataFrame:
    summary = residence_time_summary(particles)
    rows = [
        {"parametre": "Nombre total de particules", "valeur": summary["n_total"]},
        {"parametre": "Particules sorties", "valeur": summary["n_completed"]},
        {"parametre": "Particules actives (non sorties)", "valeur": summary["n_active"]},
        {"parametre": "Taux de complétion", "valeur": summary["completion_rate"]},
        {"parametre": "Temps de résidence moyen (min)", "valeur": summary["mean_residence_time_min"]},
        {"parametre": "Variance (min^2)", "valeur": summary["variance_min2"]},
        {"parametre": "Écart-type (min)", "valeur": summary["std_dev_min"]},
    ]
    return pd.DataFrame(rows)


def _group_summary_to_dataframe(particles: List[Particle]) -> pd.DataFrame:
    grouped = mean_residence_time_by_group(particles)
    rows = []
    for label, stats in sorted(grouped.items()):
        rows.append({
            "type_particule": label,
            "n_total": stats["n_total"],
            "n_sorties": stats["n_completed"],
            "n_restantes": stats["n_total"] - stats["n_completed"],
            "pourcentage_restantes": (stats["n_total"] - stats["n_completed"]) / stats["n_total"] * 100.0 if stats["n_total"] > 0 else None,
            "temps_residence_moyen_min": stats["mean_residence_time_min"],
        })
    return pd.DataFrame(rows)


def _distribution_dataframes(particles: List[Particle], n_bins: int = 30) -> Dict[str, pd.DataFrame]:
    taus = collect_residence_times(particles)

    bin_starts, counts = compute_exit_count_histogram(taus, n_bins=n_bins)
    df_hist = pd.DataFrame({"debut_intervalle_min": bin_starts, "nombre_sorties": counts})

    t_cumul, cumul_counts = compute_cumulative_exit_counts_excel(taus)
    df_cumul = pd.DataFrame({"temps_min": t_cumul, "nombre_cumule_sorties": cumul_counts})

    active_counts = count_active_particles_by_reactor(particles)
    df_active = pd.DataFrame(list(active_counts.items()), columns=["Réacteur", "Nombre de particules actives"])

    return {
        "Histogramme_sorties": df_hist,
        "Sorties_cumulees": df_cumul,
        "Particules_actives": df_active
    }


def _volume_history_to_dataframe(volume_history: Dict[str, List[float]]) -> pd.DataFrame:
    return pd.DataFrame(volume_history)


# Export principal
def export_to_excel(filepath: str, particles: List[Particle],
                    volume_history: Optional[Dict[str, List[float]]] = None,
                    simulation_config: Optional[dict] = None,
                    n_bins: int = 20) -> None:
    
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        _particles_to_dataframe(particles).to_excel(writer, sheet_name="Particules", index=False)
        _summary_to_dataframe(particles).to_excel(writer, sheet_name="Resume", index=False)
        _group_summary_to_dataframe(particles).to_excel(writer, sheet_name="Comparaison_taille", index=False)

        for name, df in _distribution_dataframes(particles, n_bins=n_bins).items():
            df.to_excel(writer, sheet_name=name, index=False)

        if volume_history:
            _volume_history_to_dataframe(volume_history).to_excel(writer, sheet_name="Volumes", index=False)

        if simulation_config:
            pd.DataFrame(
                [{"parametre": k, "valeur": v} for k, v in simulation_config.items()]
            ).to_excel(writer, sheet_name="Parametres", index=False)

    _apply_formatting_and_charts(filepath, has_volumes=bool(volume_history))


# Mise en forme + graphiques Excel
def _style_sheet(ws) -> None:
    """Police Arial, en-têtes en gras, largeur de colonnes ajustée."""
    header_font = Font(name="Arial", bold=True)
    body_font = Font(name="Arial")

    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font

    for col_cells in ws.columns:
        values = [c.value for c in col_cells if c.value is not None]
        max_len = max((len(str(v)) for v in values), default=10)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)


def _add_line_chart(ws, title: str, x_title: str, y_title: str,
                    n_rows: int, data_col: int, cat_col: int = 1,
                    anchor: str = "F2") -> None:
    if n_rows < 2:
        return
    chart = LineChart()
    chart.title = title
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    data = Reference(ws, min_col=data_col, max_col=data_col, min_row=1, max_row=n_rows)
    cats = Reference(ws, min_col=cat_col, min_row=2, max_row=n_rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def _add_bar_chart(ws, title: str, x_title: str, y_title: str,
                   n_rows: int, data_col: int, cat_col: int = 1,
                   anchor: str = "F2") -> None:
    if n_rows < 2:
        return
    chart = BarChart()
    chart.title = title
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    data = Reference(ws, min_col=data_col, max_col=data_col, min_row=1, max_row=n_rows)
    cats = Reference(ws, min_col=cat_col, min_row=2, max_row=n_rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def _apply_formatting_and_charts(filepath: str, has_volumes: bool) -> None:
    wb = load_workbook(filepath)

    # Force la première feuille à être active et visible
    if wb.worksheets:
        wb.active = 0
        for sheet in wb.worksheets:
            sheet.sheet_state = "visible"

    for ws in wb.worksheets:
        _style_sheet(ws)

    if "Histogramme_sorties" in wb.sheetnames:
        ws = wb["Histogramme_sorties"]
        _add_bar_chart(ws, "Particules sorties par intervalle", "Temps (min)", "Nombre de particules",
                        n_rows=ws.max_row, data_col=2)

    if "Sorties_cumulees" in wb.sheetnames:
        ws = wb["Sorties_cumulees"]
        _add_line_chart(ws, "Sorties cumulées du système", "Temps (min)", "Nombre cumulé",
                         n_rows=ws.max_row, data_col=2)

    if "Particules_actives" in wb.sheetnames:
        ws = wb["Particules_actives"]
        _add_bar_chart(ws, "Particules actives par réacteur", "Réacteur", "Nombre de particules actives",
                        n_rows=ws.max_row, data_col=2)

    if has_volumes and "Volumes" in wb.sheetnames:
        ws = wb["Volumes"]
        if ws.max_row >= 2:
            chart = LineChart()
            chart.title = "Suivi des volumes du système"
            chart.x_axis.title = "Temps (min)"
            chart.y_axis.title = "Volume (mL)"
            n_rows = ws.max_row
            n_cols = ws.max_column
            data = Reference(ws, min_col=2, max_col=n_cols, min_row=1, max_row=n_rows)
            cats = Reference(ws, min_col=1, min_row=2, max_row=n_rows)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, "H2")

    wb.save(filepath)