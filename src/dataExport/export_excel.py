"""
    Auteur : Lucas Durand
    Fichier d'export des résultats de simulation vers un fichier Excel, avec mise en forme et graphiques.
"""

from __future__ import annotations
from typing import List, Dict, Optional
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from openpyxl.chart.text import RichText

from simulation.particle_motion import Particle
from simulation.rtd import (
    residence_time_summary,
    collect_residence_times,
    compute_exit_count_histogram,
    compute_cumulative_exit_counts_excel,
    mean_residence_time_by_group,
    count_active_particles_by_reactor,
)
from simulation.rtd import group_particles_by_type, collect_residence_times, compute_cumulative_exit_counts_excel

COLOR_PALETTE = [
    "1F77B4", "FF7F0E", "2CA02C", "D62728", "9467BD",
    "8C564B", "E377C2", "7F7F7F", "BCBD22", "17BECF"
]

# Construction des DataFrames (une fonction par feuille)
# Construction de la DataFrame des particules 
def _particles_to_dataframe(particles: List[Particle]) -> pd.DataFrame:
    rows = []
    for i, p in enumerate(particles):
        rows.append({
            "id": i + 1,
            "densite_kg_m3": p.particle_type.particle_density,
            "rayon_mm": p.particle_type.particle_size * 1000.0,
            "reacteur_final": p.current_reactor,
            "temps_entree_min": p.entry_time_s / 60.0 if p.entry_time_s is not None else None,
            "temps_sortie_min": p.exit_time_s / 60.0 if p.exit_time_s is not None else None,
            "temps_residence_min": p.residence_time_min,
            "est_sortie": p.residence_time_min is not None,
        })
    return pd.DataFrame(rows)

# COnstruction de la DataFrame du résumé des résultats
def _summary_to_dataframe(particles: List[Particle]) -> pd.DataFrame:
    summary = residence_time_summary(particles)
    rows = [
        {"parametre": "Nombre total de particules", "valeur": summary["n_total"]},
        {"parametre": "Particules sorties", "valeur": summary["n_completed"]},
        {"parametre": "Particules actives (non sorties)", "valeur": summary["n_active"]},
        {"parametre": "Taux de complétion", "valeur": summary["completion_rate"]},
        {"parametre": "Temps de résidence moyen (min)", "valeur": summary["mean_residence_time_min"]},
        {"parametre": "Variance (min^2)", "valeur": summary["variance_min2"]},
        {"parametre": "Écart-type (min)", "valeur": summary["std_dev_min"]},
    ]
    return pd.DataFrame(rows)

# Construction de la DataFrame du résumé par groupe (densité/taille)
def _group_summary_to_dataframe(particles: List[Particle]) -> pd.DataFrame:
    grouped = mean_residence_time_by_group(particles)
    rows = []
    for label, stats in sorted(grouped.items()):
        rows.append({
            "type_particule": label,
            "n_total": stats["n_total"],
            "n_sorties": stats["n_completed"],
            "n_restantes": stats["n_total"] - stats["n_completed"],
            "pourcentage_restantes": (stats["n_total"] - stats["n_completed"]) / stats["n_total"] * 100.0 if stats["n_total"] > 0 else None,
            "temps_residence_moyen_min": stats["mean_residence_time_min"],
        })
    return pd.DataFrame(rows)

#
def _cumulative_by_group_to_dataframe(particles: List[Particle]) -> pd.DataFrame:
    """
    Génère un DataFrame avec les courbes de sorties cumulées pour chaque groupe (densité/taille).
    """

    groups = group_particles_by_type(particles)
    data = {}

    for label, group_particles in groups.items():
        taus = collect_residence_times(group_particles)
        t_cumul, cumul_counts = compute_cumulative_exit_counts_excel(taus)
        
        # On crée 2 colonnes par groupe : Temps et Nombre cumulé
        data[f"Temps_{label}"] = t_cumul
        data[f"Cumul_{label}"] = cumul_counts

    # On aligne les colonnes de longueurs différentes avec du vide (NaN)
    df = pd.DataFrame(dict([(k, pd.Series(v)) for k, v in data.items()]))
    return df

# Construction des DataFrames pour les distributions et les histogrammes
def _distribution_dataframes(particles: List[Particle], n_bins: int = 30) -> Dict[str, pd.DataFrame]:
    taus = collect_residence_times(particles)

    bin_starts, counts = compute_exit_count_histogram(taus, n_bins=n_bins)
    df_hist = pd.DataFrame({"debut_intervalle_min": bin_starts, "nombre_sorties": counts})

    t_cumul, cumul_counts = compute_cumulative_exit_counts_excel(taus)
    df_cumul = pd.DataFrame({"temps_min": t_cumul, "nombre_cumule_sorties": cumul_counts})

    active_counts = count_active_particles_by_reactor(particles)
    df_active = pd.DataFrame(list(active_counts.items()), columns=["Réacteur", "Nombre de particules actives"])

    return {
        "Histogramme_sorties": df_hist,
        "Sorties_cumulees": df_cumul,
        "Particules_actives": df_active,
    }

#Construction de la DataFrame du suivi des volumes 
def _volume_history_to_dataframe(volume_history: Dict[str, List[float]]) -> pd.DataFrame:
    df = pd.DataFrame(volume_history)
    
    if "t" in df.columns:
        df["temps_min"] = df["t"] / 60.0
        df = df.drop(columns=["t"])  # supprime la colonne originale (temps en secondes)

        #Placement de la colonne "temps_min" en première position
        other_cols = [col for col in df.columns if col != 'temps_min']
        df = df[['temps_min'] + other_cols]
        
    return df


# Export principal des Dataframes vers le fichier excel, avec mise en forme et graphiques
def export_to_excel(filepath: str, particles: List[Particle],
                    volume_history: Optional[Dict[str, List[float]]] = None,
                    simulation_config: Optional[dict] = None,
                    n_bins: int = 20) -> None:
    
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        _particles_to_dataframe(particles).to_excel(writer, sheet_name="Particules", index=False)
        _summary_to_dataframe(particles).to_excel(writer, sheet_name="Resume", index=False)
        _group_summary_to_dataframe(particles).to_excel(writer, sheet_name="Comparaison_taille", index=False)
        _cumulative_by_group_to_dataframe(particles).to_excel(writer, sheet_name="Cumul_par_groupe", index=False)

        for name, df in _distribution_dataframes(particles, n_bins=n_bins).items():
            df.to_excel(writer, sheet_name=name, index=False)


        if volume_history:
            _volume_history_to_dataframe(volume_history).to_excel(writer, sheet_name="Volumes", index=False)

        if simulation_config:
            pd.DataFrame(
                [{"parametre": k, "valeur": v} for k, v in simulation_config.items()]
            ).to_excel(writer, sheet_name="Parametres", index=False)

    _apply_formatting_and_charts(filepath, has_volumes=bool(volume_history))


# Mise en forme + graphiques Excel
def _style_sheet(ws) -> None:
    """Police Arial, en-têtes en gras, largeur de colonnes ajustée."""
    header_font = Font(name="Arial", bold=True)
    body_font = Font(name="Arial")

    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font

    for col_cells in ws.columns:
        values = [c.value for c in col_cells if c.value is not None]
        max_len = max((len(str(v)) for v in values), default=10)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

def _create_base_chart(chart_cls, title: str, x_title: str, y_title: str, style: int | None = None):
    """Crée et initialise les propriétés communes d'un graphique OpenPyXL."""
    chart = chart_cls()
    chart.title = title
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title

    chart.x_axis.title = x_title
    chart.y_axis.title = y_title

    chart.x_axis.delete = False
    chart.y_axis.delete = False

    chart.x_axis.tickLblPos = "low"

    chart.width = 22
    chart.height = 11

    _shrink_axis_labels(chart.x_axis, size=800)
    _shrink_axis_labels(chart.y_axis, size=800)

    _position_axis_title(chart.x_axis, x=0.35, y=0.92)
    _position_axis_title(chart.y_axis, x=0.02, y=0.35)

    if style is not None:
        chart.style = style
    return chart

def _position_axis_title(axis, x, y):
    """Positionne manuellement le titre d'un axe (coordonnées relatives 0-1)."""
    if axis.title is None:
        return
    axis.title.layout = Layout(
        manualLayout=ManualLayout(
            xMode="edge",
            yMode="edge",
            x=x,
            y=y,
        )
    )

def _shrink_axis_labels(axis, size=800):
    """Réduit la taille de police des labels d'un axe (size en centièmes de point, 800 = 8pt)."""
    axis.txPr = RichText(
        bodyPr=None,
        p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=size)))]
    )

def _style_series_line(series, color_hex: str, width: int = 25000) -> None:
    """Applique une couleur et une épaisseur à une série de lignes."""
    series.graphicalProperties.line.solidFill = color_hex
    series.graphicalProperties.line.width = width

# Graphiques simples (Ligne ou Barre mono-série)
def add_single_series_chart(ws, chart_type: str, title: str, x_title: str, y_title: str,
                             n_rows: int, data_col: int, cat_col: int = 1,
                             anchor: str = "F2", style: int | None = 13) -> None:
    """Génère un graphique à une seule série de données (Line ou Bar)"""
    if n_rows < 2:
        return

    chart_cls = LineChart if chart_type == "line" else BarChart
    chart = _create_base_chart(chart_cls, title, x_title, y_title, style=style)

    data = Reference(ws, min_col=data_col, max_col=data_col, min_row=1, max_row=n_rows)
    cats = Reference(ws, min_col=cat_col, min_row=2, max_row=n_rows)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)

# Alias pratiques pour garder une compatibilité avec ton code existant
def _add_line_chart(ws, title: str, x_title: str, y_title: str, n_rows: int, data_col: int, cat_col: int = 1, anchor: str = "F2") -> None:
    add_single_series_chart(ws, "line", title, x_title, y_title, n_rows, data_col, cat_col, anchor, style=13)

def _add_bar_chart(ws, title: str, x_title: str, y_title: str, n_rows: int, data_col: int, cat_col: int = 1, anchor: str = "F2") -> None:
    add_single_series_chart(ws, "bar", title, x_title, y_title, n_rows, data_col, cat_col, anchor, style=13)

def _add_bar_chart_active(ws, title: str, x_title: str, y_title: str, n_rows: int, data_col: int, cat_col: int = 1, anchor: str = "F2") -> None:
    add_single_series_chart(ws, "bar", title, x_title, y_title, n_rows, data_col, cat_col, anchor, style=None)


# Graphiques multi-lignes

def _add_multi_line_chart(ws, title: str, x_title: str, y_title: str, anchor: str = "F2") -> None:
    """Graphique multi-lignes pour colonnes associées par paires (Temps, Cumul)."""
    if ws.max_row < 2 or ws.max_column < 2:
        return

    chart = _create_base_chart(LineChart, title, x_title, y_title)
    n_rows, n_cols = ws.max_row, ws.max_column

    for idx, col in enumerate(range(1, n_cols, 2)):
        time_col, data_col = col, col + 1

        data = Reference(ws, min_col=data_col, max_col=data_col, min_row=1, max_row=n_rows)
        cats = Reference(ws, min_col=time_col, min_row=2, max_row=n_rows)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        _style_series_line(chart.series[-1], COLOR_PALETTE[idx % len(COLOR_PALETTE)])

    ws.add_chart(chart, anchor)

def _add_multi_line_chart_volume(ws, title: str, x_title: str, y_title: str, anchor: str = "F2") -> None:
    """Graphique multi-lignes avec un axe X commun (colonne 1) et N séries de données."""
    if ws.max_row < 2 or ws.max_column < 2:
        return

    chart = _create_base_chart(LineChart, title, x_title, y_title)
    
    cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    data = Reference(ws, min_col=2, max_col=ws.max_column, min_row=1, max_row=ws.max_row)

    chart.set_categories(cats)
    chart.add_data(data, titles_from_data=True)

    for idx, series in enumerate(chart.series):
        _style_series_line(series, COLOR_PALETTE[idx % len(COLOR_PALETTE)])

    ws.add_chart(chart, anchor)

# Applique la mise en forme et les graphiques aux bonnes données et sur les bonnes feuilles du fichier Excel
def _apply_formatting_and_charts(filepath: str, has_volumes: bool) -> None:
    wb = load_workbook(filepath)

    # Force la première feuille à être active et visible
    if wb.worksheets:
        wb.active = 0
        for sheet in wb.worksheets:
            sheet.sheet_state = "visible"

    for ws in wb.worksheets:
        _style_sheet(ws)

    if "Histogramme_sorties" in wb.sheetnames:
        ws = wb["Histogramme_sorties"]
        _add_bar_chart(ws, "Particules sorties par intervalle", "Temps (min)", "Nombre de particules",
                        n_rows=ws.max_row, data_col=2)

    if "Sorties_cumulees" in wb.sheetnames:
        ws = wb["Sorties_cumulees"]
        _add_line_chart(ws, "Sorties cumulées du système", "Temps (min)", "Nombre cumulé",
                         n_rows=ws.max_row, data_col=2)

    if "Cumul_par_groupe" in wb.sheetnames:
        ws = wb["Cumul_par_groupe"]
        _add_multi_line_chart(
            ws, 
            title="Sorties cumulées par type de particule (taille/densité)", 
            x_title="Temps (min)", 
            y_title="Nombre cumulé de sorties", 
            anchor="F2"
        )

    if "Particules_actives" in wb.sheetnames:
        ws = wb["Particules_actives"]
        _add_bar_chart_active(ws, "Particules actives par réacteur", "Réacteur", "Nombre de particules actives",
                        n_rows=ws.max_row, data_col=2)

    if has_volumes and "Volumes" in wb.sheetnames:
        ws = wb["Volumes"]
        if ws.max_row >= 2:
            _add_multi_line_chart_volume(
                ws,
                title="Suivi des volumes du système",
                x_title="Temps (min)",
                y_title="Volume (mL)",
                anchor="H2"
            )
            
    wb.save(filepath)